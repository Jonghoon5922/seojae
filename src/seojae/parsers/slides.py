"""PowerPoint(.pptx) 파서.

슬라이드 하나가 덩어리 하나다. 출처는 "슬라이드 7 · 전환 일정" 처럼 남긴다.
발표자 노트도 읽는다 — 슬라이드 본문은 키워드만 있고 설명은 노트에 있는 일이 흔하다.
"""

from __future__ import annotations

from pathlib import Path

from .base import Block, ParsedDoc, ParseError, register


def _shape_text(shape) -> str:
    """도형 하나의 글. 표는 행 단위로 편다."""
    if shape.has_table:
        rows = []
        for row in shape.table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        return "\n".join(rows)

    if shape.has_text_frame:
        return shape.text_frame.text.strip()

    return ""


@register(".pptx", label="PowerPoint 문서")
def parse_pptx(path: Path) -> ParsedDoc:
    try:
        from pptx import Presentation

        deck = Presentation(str(path))
    except Exception as e:
        raise ParseError(f"pptx를 열지 못했다: {e}") from e

    blocks: list[Block] = []
    title = path.stem

    for index, slide in enumerate(deck.slides, start=1):
        pieces: list[str] = []
        heading = ""

        for shape in slide.shapes:
            try:
                text = _shape_text(shape)
            except Exception:
                continue  # 그림·차트 등 글이 없는 도형
            if not text:
                continue
            if not heading:
                heading = text.splitlines()[0][:60]
            pieces.append(text)

        # 발표자 노트
        try:
            if slide.has_notes_slide:
                note = slide.notes_slide.notes_text_frame.text.strip()
                if note:
                    pieces.append(f"[노트] {note}")
        except Exception:
            pass

        if not pieces:
            continue

        if index == 1 and heading and title == path.stem:
            title = heading

        location = f"슬라이드 {index}" + (f" · {heading}" if heading else "")
        blocks.append(Block(text="\n".join(pieces), location=location, group=location))

    if not blocks:
        raise ParseError("pptx에서 읽어낸 글이 없다.")

    return ParsedDoc(title=title, blocks=blocks, pages=len(deck.slides))
